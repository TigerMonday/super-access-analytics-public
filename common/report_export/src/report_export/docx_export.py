"""MarkdownからWord(docx)への変換。

方針:
- pandocが使える環境ではpandocに委譲する(表・見出し・強調・日本語の再現度が高い)。
- pandocが無い環境ではpython-docxで簡易変換する(見出し・段落・箇条書き・表・
  太字/イタリック/インラインコードのみ対応。装飾はpandoc版より簡素)。
"""

from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from .md_tables import _HEADING_RE, _SEP_LINE_RE, parse_markdown_for_excel

_INLINE_TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|_[^_]+_|`[^`]+`)")
_BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUM_RE = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")


def pandoc_available() -> bool:
    return shutil.which("pandoc") is not None


def export_docx_via_pandoc(md_path: Path, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        "pandoc",
        str(md_path),
        "-f",
        "gfm",
        "-o",
        str(output_path),
        "--resource-path",
        str(md_path.parent),
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(f"pandocによるdocx変換に失敗しました: {result.stderr}")
    return output_path


def _add_inline_runs(paragraph, text: str) -> None:
    parts = _INLINE_TOKEN_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Menlo"
        elif (part.startswith("*") and part.endswith("*") and len(part) > 2) or (
            part.startswith("_") and part.endswith("_") and len(part) > 2
        ):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def export_docx_via_python_docx(md_path: Path, output_path: Path, md_text: str) -> Path:
    # python-docxは呼び出し時のみ遅延importする(pandoc経路では不要な依存のため)。
    import docx  # type: ignore

    doc = docx.Document()
    parsed = parse_markdown_for_excel(md_text)
    table_by_start = {t.start_line: t for t in parsed.tables}

    lines = md_text.splitlines()
    n = len(lines)
    i = 0
    in_code_fence = False
    code_buffer: list[str] = []

    def flush_code() -> None:
        nonlocal code_buffer
        if code_buffer:
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_buffer))
            run.font.name = "Menlo"
            code_buffer = []

    while i < n:
        line = lines[i]

        if i in table_by_start:
            table = table_by_start[i]
            n_cols = max(len(table.headers), max((len(r) for r in table.rows), default=0)) or 1
            docx_table = doc.add_table(rows=1, cols=n_cols)
            # アクセント無しの中立なグリッド(design-system.mdのグレースケース基調に合わせる)
            docx_table.style = "Light Grid"
            hdr_cells = docx_table.rows[0].cells
            for c, val in enumerate(table.headers):
                hdr_cells[c].text = val
            for row in table.rows:
                cells = docx_table.add_row().cells
                for c, val in enumerate(row):
                    if c < n_cols:
                        cells[c].text = val
            doc.add_paragraph("")
            i = table.end_line
            continue

        if line.strip().startswith("```") or line.strip().startswith("~~~"):
            in_code_fence = not in_code_fence
            if not in_code_fence:
                flush_code()
            i += 1
            continue

        if in_code_fence:
            code_buffer.append(line)
            i += 1
            continue

        h = _HEADING_RE.match(line)
        if h:
            level = min(len(h.group(1)), 4)
            doc.add_heading(h.group(2).strip(), level=level)
            i += 1
            continue

        if _SEP_LINE_RE.match(line.strip()) and line.strip():
            i += 1
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(p, bullet.group(2))
            i += 1
            continue

        numbered = _NUM_RE.match(line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            _add_inline_runs(p, numbered.group(2))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        p = doc.add_paragraph()
        _add_inline_runs(p, line.strip())
        i += 1

    flush_code()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def export_docx(md_path: Path, output_path: Path, md_text: str) -> Path:
    if pandoc_available():
        return export_docx_via_pandoc(md_path, output_path)
    return export_docx_via_python_docx(md_path, output_path, md_text)
