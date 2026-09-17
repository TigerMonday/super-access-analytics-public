"""fixturesを実際に4形式へ変換し、壊れていないことを確認する統合テスト。"""

from pathlib import Path
import re

import docx as docx_lib
import pytest
from openpyxl import load_workbook

from report_export.__main__ import run

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "fixtures"

FIXTURE_FILES = [
    FIXTURES_DIR / "with_tables.md",
    FIXTURES_DIR / "no_tables.md",
    FIXTURES_DIR / "with_chart.md",
]


@pytest.mark.parametrize("md_path", FIXTURE_FILES, ids=lambda p: p.stem)
def test_all_formats_generated_and_openable(md_path: Path, tmp_path: Path):
    # PDFの実起動は画像・表を含む代表入力だけ。他形式の境界は全入力で維持する。
    formats = ["html", "docx", "xlsx"]
    if md_path.stem == "with_chart":
        formats.append("pdf")
    generated = run(md_path, formats, tmp_path)

    by_suffix = {p.suffix: p for p in generated}
    assert set(by_suffix) == {f".{fmt}" for fmt in formats}

    for path in generated:
        assert path.is_file()
        assert path.is_absolute()
        assert path.stat().st_size > 0

    # HTML: 単体完結の構造チェック
    html_text = by_suffix[".html"].read_text(encoding="utf-8")
    assert "<html" in html_text
    assert "</html>" in html_text
    assert "<table" in html_text or "no_tables" in md_path.stem

    if ".pdf" in by_suffix:
        pdf_bytes = by_suffix[".pdf"].read_bytes()
        assert pdf_bytes.startswith(b"%PDF")
        boxes = re.findall(
            rb"/MediaBox\s*\[\s*([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*\]",
            pdf_bytes,
        )
        assert boxes, "PDFにページ寸法がない"
        for box in boxes:
            x0, y0, x1, y1 = map(float, box)
            # 1280×720 CSS px = 13.333333×7.5 inch（PDF座標は72pt/in）。
            assert x1 - x0 == pytest.approx(13.333333 * 72, abs=1)
            assert y1 - y0 == pytest.approx(7.5 * 72, abs=1)
        assert b"/Subtype /Image" in pdf_bytes, "元のチャート画像がPDFから欠落"

    # docx: python-docxで再読込できる(壊れていない)
    document = docx_lib.Document(str(by_suffix[".docx"]))
    assert len(document.paragraphs) > 0

    # xlsx: openpyxlで再読込できる(壊れていない)
    wb = load_workbook(by_suffix[".xlsx"])
    assert "概要" in wb.sheetnames


def test_no_tables_fixture_xlsx_has_only_overview_sheet(tmp_path: Path):
    md_path = FIXTURES_DIR / "no_tables.md"
    generated = run(md_path, ["xlsx"], tmp_path)
    xlsx_path = generated[0]
    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == ["概要"]


def test_with_tables_fixture_xlsx_has_one_sheet_per_table(tmp_path: Path):
    md_path = FIXTURES_DIR / "with_tables.md"
    generated = run(md_path, ["xlsx"], tmp_path)
    wb = load_workbook(generated[0])
    # 概要 + チャネル別サマリー + 施策別メモ(改行・空セルあり) の2表 = 3シート
    assert len(wb.sheetnames) == 3
    assert "概要" in wb.sheetnames


def test_with_tables_fixture_xlsx_escaped_pipe_and_br_survive(tmp_path: Path):
    md_path = FIXTURES_DIR / "with_tables.md"
    generated = run(md_path, ["xlsx"], tmp_path)
    wb = load_workbook(generated[0])
    channel_sheet = wb["チャネル別サマリー"]
    rows = list(channel_sheet.iter_rows(values_only=True))
    flat = [cell for row in rows for cell in row if cell]
    assert any("Referral|提携" == v for v in flat)

    memo_sheet_name = [n for n in wb.sheetnames if n.startswith("施策別メモ")][0]
    memo_sheet = channel_sheet = wb[memo_sheet_name]
    rows = list(memo_sheet.iter_rows(values_only=True))
    flat_text = " ".join(str(c) for row in rows for c in row if c)
    assert "A/Bテスト中\n" in flat_text or "A/Bテスト中" in flat_text


def test_with_chart_fixture_html_embeds_base64_image(tmp_path: Path):
    md_path = FIXTURES_DIR / "with_chart.md"
    generated = run(md_path, ["html"], tmp_path)
    html_text = generated[0].read_text(encoding="utf-8")
    assert "data:image/png;base64," in html_text
    # 元の相対パス参照は残っていない
    assert "charts/sample_chart.png" not in html_text


def test_html_title_is_not_duplicated_in_body(tmp_path: Path):
    # <title>とheaderのh1には出るが、本文側では元のH1を重複表示しない。
    # 本文はセクションごとの白い面(<section class="report-sheet report-body">)として
    # 表紙の面のあとに置かれる(report-template.html参照)ため、そのタグを境目に
    # 本文側(表紙を除く)だけを取り出して確認する。
    md_path = FIXTURES_DIR / "with_tables.md"
    generated = run(md_path, ["html"], tmp_path)
    html_text = generated[0].read_text(encoding="utf-8")
    title = "集客レポートサンプル（表あり）"
    body = html_text.split('<section class="report-sheet report-body">', 1)[1]
    assert title not in body


def test_docx_python_docx_fallback_path_works_when_pandoc_absent(tmp_path: Path):
    """pandoc未導入環境向けfallback(python-docx)経路も壊れていないことを確認する。"""
    from report_export.docx_export import export_docx_via_python_docx
    from report_export.markdown_utils import read_markdown

    md_path = FIXTURES_DIR / "with_tables.md"
    md_text = read_markdown(md_path)
    out_path = tmp_path / "fallback.docx"
    export_docx_via_python_docx(md_path, out_path, md_text)

    assert out_path.stat().st_size > 0
    document = docx_lib.Document(str(out_path))
    assert len(document.paragraphs) > 0
    assert len(document.tables) == 2
    header_row_texts = [c.text for c in document.tables[0].rows[0].cells]
    assert header_row_texts == ["チャネル", "セッション数", "CV数", "CVR"]


def test_docx_export_dispatches_to_pandoc_when_available(tmp_path: Path, monkeypatch):
    from report_export import docx_export

    calls = {"pandoc": 0, "fallback": 0}

    def fake_pandoc(md_path, output_path):
        calls["pandoc"] += 1
        output_path.write_bytes(b"dummy")
        return output_path

    def fake_fallback(md_path, output_path, md_text):
        calls["fallback"] += 1
        output_path.write_bytes(b"dummy")
        return output_path

    monkeypatch.setattr(docx_export, "pandoc_available", lambda: True)
    monkeypatch.setattr(docx_export, "export_docx_via_pandoc", fake_pandoc)
    monkeypatch.setattr(docx_export, "export_docx_via_python_docx", fake_fallback)

    docx_export.export_docx(FIXTURES_DIR / "with_tables.md", tmp_path / "x.docx", "dummy")
    assert calls == {"pandoc": 1, "fallback": 0}

    monkeypatch.setattr(docx_export, "pandoc_available", lambda: False)
    docx_export.export_docx(FIXTURES_DIR / "with_tables.md", tmp_path / "y.docx", "dummy")
    assert calls == {"pandoc": 1, "fallback": 1}


def test_output_dir_option_is_respected(tmp_path: Path):
    md_path = FIXTURES_DIR / "no_tables.md"
    custom_dir = tmp_path / "custom_out"
    generated = run(md_path, ["html"], custom_dir)
    assert generated[0].parent == custom_dir.resolve()


def test_accent_hex_none_keeps_default_gold_tokens(tmp_path: Path):
    """accent_hex未指定なら、既定のゴールド(#D4AF37系)のまま変わらない。"""
    md_path = FIXTURES_DIR / "no_tables.md"
    generated = run(md_path, ["html"], tmp_path)
    html_text = generated[0].read_text(encoding="utf-8")
    assert "--gold:#D4AF37;" in html_text
    # <style>内(実際に効くCSS)には上書き用の:rootブロックが差し込まれない
    # ("\n<style>\n"で分割する。先頭コメントの説明文中にも"<style>"という語が出るため単純split(<style>)では誤爆する)
    style_text = html_text.split("\n<style>\n", 1)[1].split("</style>", 1)[0]
    assert style_text.count(":root{") == 1


def test_accent_hex_overrides_gold_tokens_via_second_root_block(tmp_path: Path):
    md_path = FIXTURES_DIR / "no_tables.md"
    generated = run(md_path, ["html"], tmp_path, accent_hex="#1D4ED8")
    html_text = generated[0].read_text(encoding="utf-8")
    # 元のコアCSSのゴールド定義は書き換えず残したまま、
    assert "--gold:#D4AF37;" in html_text
    # 後勝ちの2つ目の:rootブロックで上書きする(<style>内で数える。先頭コメントの説明文と混同しない)
    style_text = html_text.split("\n<style>\n", 1)[1].split("</style>", 1)[0]
    assert style_text.count(":root{") == 2
    assert "--gold:#1d4ed8;" in html_text
    assert "--gold-dark:#1d4ed8;" in html_text


def test_accent_hex_gold_dark_is_readable_even_for_bright_input(tmp_path: Path):
    """derive_paletteのコントラスト保証がCLI経由の実出力にも効いていることを確認する。"""
    from report_export.accent_color import MIN_TEXT_CONTRAST, contrast_ratio, hex_to_rgb

    md_path = FIXTURES_DIR / "no_tables.md"
    generated = run(md_path, ["html"], tmp_path, accent_hex="#FFC800")
    html_text = generated[0].read_text(encoding="utf-8")
    line = [ln for ln in html_text.splitlines() if ln.strip().startswith("--gold-dark:")][-1]
    gold_dark_hex = line.split(":", 1)[1].strip().rstrip(";")
    ratio = contrast_ratio(hex_to_rgb(gold_dark_hex), (255, 255, 255))
    assert ratio >= MIN_TEXT_CONTRAST


def test_logo_path_none_shows_no_logo(tmp_path: Path):
    """--logo未指定(既定)では、ロゴのCSSも表紙の要素(固定サイズの空箱)も出ない。"""
    md_path = FIXTURES_DIR / "no_tables.md"
    generated = run(md_path, ["html"], tmp_path)
    html_text = generated[0].read_text(encoding="utf-8")

    assert "cover-logo" not in html_text
    assert "data:image/svg+xml;base64," not in html_text


def test_logo_path_given_shows_logo(tmp_path: Path):
    """--logo(logo_path)を指定したときだけ、data URIとロゴ要素が出る。"""
    from report_export import house_style

    custom_logo = tmp_path / "custom-logo.svg"
    custom_logo.write_text('<svg xmlns="http://www.w3.org/2000/svg"></svg>', encoding="utf-8")

    md_path = FIXTURES_DIR / "no_tables.md"
    generated = run(md_path, ["html"], tmp_path, logo_path=custom_logo)
    html_text = generated[0].read_text(encoding="utf-8")

    assert 'class="cover-logo"' in html_text
    assert house_style._logo_base64(custom_logo) in html_text
