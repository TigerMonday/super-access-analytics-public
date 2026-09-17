"""xlsxシート名生成ロジックのテスト。

Excelのシート名制約(31文字上限、禁止文字)を満たしつつ、
長い同一見出しに由来する複数の表が区別可能な名前になることを検証する。
"""

from pathlib import Path

from openpyxl import load_workbook

import pytest

from report_export.xlsx_export import (
    MAX_SHEET_NAME_LEN,
    export_xlsx,
    generate_sheet_name,
)


@pytest.mark.parametrize("value", ["=1+2", "+1+2", "-1+2", "@SUM(A1)"])
def test_untrusted_values_remain_text_after_save(tmp_path: Path, value: str):
    md = f"# {value}\n\n| {value} |\n|---|\n| {value} |\n"
    out = tmp_path / "safe.xlsx"
    export_xlsx(tmp_path / "source.md", out, md)
    wb = load_workbook(out, data_only=False)
    cells = [wb.worksheets[0]["A3"], wb.worksheets[1]["A1"], wb.worksheets[1]["A2"]]
    for cell in cells:
        assert cell.value == value
        assert cell.data_type == "s"
        assert cell.quotePrefix is True


def test_short_name_is_unchanged():
    used: set[str] = set()
    assert generate_sheet_name("チャネル別サマリー", used) == "チャネル別サマリー"


def test_forbidden_characters_are_removed():
    used: set[str] = set()
    name = generate_sheet_name("5:流入質/評価[参考]", used)
    assert len(name) <= MAX_SHEET_NAME_LEN
    for ch in "\\/?*[]:":
        assert ch not in name


def test_long_heading_is_shortened_within_limit_and_keeps_prefix():
    used: set[str] = set()
    long_heading = "3. 媒体別パフォーマンス（utm_source × utm_medium）"
    assert len(long_heading) > MAX_SHEET_NAME_LEN

    name = generate_sheet_name(long_heading, used)
    assert len(name) <= MAX_SHEET_NAME_LEN
    # 番号プレフィックスは短縮後も残る
    assert name.startswith("3.")


def test_two_long_tables_with_same_heading_prefix_stay_distinguishable():
    """同じ見出し配下の2表(md_tables側で既に「_2」等の連番が付与済み)が
    31文字に短縮された後も、単純な前方truncateで衝突しないことを確認する。
    """
    used: set[str] = set()
    base = "3. 媒体別パフォーマンス（utm_source × utm_medium）"
    title1 = base
    title2 = f"{base}_2"

    name1 = generate_sheet_name(title1, used)
    used.add(name1.lower())
    name2 = generate_sheet_name(title2, used)

    assert len(name1) <= MAX_SHEET_NAME_LEN
    assert len(name2) <= MAX_SHEET_NAME_LEN
    assert name1 != name2
    assert name1.lower() != name2.lower()
    # 差分がわかるよう、2枚目には連番の痕跡が残っている
    assert name2 != name1 and (name2.endswith("_2") or name1 not in name2)


def test_collision_falls_back_to_parenthesized_number():
    """短縮結果が既存名と衝突する極端なケースでは、最後の手段として (2) を使う。"""
    long_name = "a" * 50
    first = generate_sheet_name(long_name, set())
    used = {first.lower()}

    second = generate_sheet_name(long_name, used)
    assert len(second) <= MAX_SHEET_NAME_LEN
    assert second != first
    assert second.endswith("(2)")


def test_no_preceding_heading_gets_default_name(tmp_path: Path):
    """見出しの前に表が出現するケース。既定名(表1等)が31文字以内かつ一意になる。"""
    md = """
| a | b |
|---|---|
| 1 | 2 |

## セクション

| c | d |
|---|---|
| 3 | 4 |
"""
    md_path = tmp_path / "no_heading.md"
    out_path = tmp_path / "out.xlsx"
    export_xlsx(md_path, out_path, md)

    wb = load_workbook(out_path)
    assert wb.sheetnames[0] == "概要"
    assert "表1" in wb.sheetnames
    assert all(len(n) <= MAX_SHEET_NAME_LEN for n in wb.sheetnames)
    assert len(wb.sheetnames) == len(set(n.lower() for n in wb.sheetnames))
